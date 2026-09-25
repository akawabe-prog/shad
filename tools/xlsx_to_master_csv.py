#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
商品マスターが xlsx で届いたときに、パイプラインが読む data-source/ItemList_SHAD.csv（cp932）へ変換する。

    python3 tools/xlsx_to_master_csv.py ~/Downloads/ItemList_260924152041_260924152109.xlsx

・1シート目を読む。列名はそのまま、数値は整数/小数の表記を CSV 出力と同じにする（33.0 → 33）
・既存の CSV は data-source/ItemList_SHAD_<日付>_backup.csv に退避
"""
import sys, os, csv, shutil, datetime
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data-source", "ItemList_SHAD.csv")

def cv(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v)

def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = sys.argv[1]
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [cv(h) for h in rows[0]]
    if os.path.exists(OUT):
        bk = OUT.replace(".csv", "_%s_backup.csv" % datetime.date.today().strftime("%y%m%d"))
        shutil.copy(OUT, bk)
        print("退避:", os.path.relpath(bk, ROOT))
    with open(OUT, "w", encoding="cp932", newline="", errors="replace") as f:
        w = csv.writer(f)
        w.writerow(hdr)
        n = 0
        for r in rows[1:]:
            if all(v is None for v in r):
                continue
            w.writerow([cv(v) for v in r]); n += 1
    print("出力: %s（%d 行 / %d 列）" % (os.path.relpath(OUT, ROOT), n, len(hdr)))
    print("次に: python3 tools/build_catalog.py → apply_master_to_pages.py → MAX再生成 → build_cards_json.py → audit_master.py")

if __name__ == "__main__":
    main()
